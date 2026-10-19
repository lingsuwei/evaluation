"""评测集 JSON / CSV 导入导出。"""
from __future__ import annotations

import io
import json

import pytest
from httpx import ASGITransport, AsyncClient
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker, create_async_engine

from database import Base, Dataset, get_session
from main import app
from services.dataset_io_service import (
    DatasetUploadFormatError,
    SCENARIO_STEP_EXPORT_HEADERS,
    _step_to_export_fields,
    content_disposition_attachment,
    csv_template_bytes,
    default_dataset_name,
    detect_upload_format,
    export_dataset_csv_bytes,
    export_dataset_json_bytes,
    export_dataset_xlsx_bytes,
    parse_csv_upload,
    parse_json_upload,
    parse_scenario_csv,
    parse_upload_content,
    validate_dataset_items_with_format,
)


def test_parse_csv_with_chinese_headers():
    content = (
        "题目,标准答案,模型答案\n"
        "Q1,GT1,A1\n"
        "Q2,GT2,A2\n"
    ).encode("utf-8-sig")
    items = parse_csv_upload(content)
    assert len(items) == 2
    assert items[0]["question"] == "Q1"
    assert items[0]["ground_truth"] == "GT1"
    assert items[0]["model_answer"] == "A1"


def test_csv_roundtrip_preserves_fields():
    original = [
        {
            "id": "ditem_abc",
            "question": "题目",
            "prompt": "提示词",
            "ground_truth": "标准",
            "model_answer": "答案",
            "judging_mode": "Ground_Truth评测",
        }
    ]
    csv_bytes = export_dataset_csv_bytes(original)
    parsed = parse_csv_upload(csv_bytes)
    assert parsed[0]["question"] == "题目"
    assert parsed[0]["id"] == "ditem_abc"
    assert parsed[0]["prompt"] == "提示词"
    assert parsed[0]["judging_mode"] == "Ground_Truth评测"


def test_validate_csv_row_errors():
    content = "题目,模型答案\n,empty\n".encode("utf-8")
    with pytest.raises(Exception):
        parse_csv_upload(content)


def test_json_upload_unchanged():
    data = [{"question": "q", "model_answer": "a"}]
    raw = json.dumps(data).encode("utf-8")
    parsed = parse_upload_content(raw, "test.json")
    assert parsed == data
    assert validate_dataset_items_with_format(parsed, "test.json") == []


# --- 格式检测 / 工具函数 ---


def test_detect_upload_format_accepts_json_and_csv():
    assert detect_upload_format("data.JSON") == "json"
    assert detect_upload_format("path/to/sample.csv") == "csv"


def test_detect_upload_format_rejects_unsupported():
    with pytest.raises(DatasetUploadFormatError, match="仅支持"):
        detect_upload_format("notes.txt")


def test_default_dataset_name_strips_extension():
    assert default_dataset_name(None, "我的集.csv", "ds_abc") == "我的集"
    assert default_dataset_name("  自定义  ", "x.json", "ds_abc") == "自定义"
    assert default_dataset_name(None, None, "ds_xyz") == "dataset_ds_xyz"


def test_content_disposition_attachment_chinese_filename():
    header = content_disposition_attachment("导出测试.csv")
    assert "filename*=" in header
    assert "UTF-8" in header
    assert "attachment" in header


def test_csv_template_has_bom_and_example_row():
    body = csv_template_bytes()
    assert body.startswith(b"\xef\xbb\xbf")
    text = body.decode("utf-8-sig")
    assert "题目" in text
    assert "示例题目" in text


# --- JSON 解析：正向 / 反向 / 异常 ---


def test_parse_json_upload_accepts_array():
    items = parse_json_upload(b'[{"question":"q","model_answer":"a"}]')
    assert len(items) == 1


def test_parse_json_upload_rejects_invalid_json():
    with pytest.raises(DatasetUploadFormatError, match="Invalid JSON"):
        parse_json_upload(b"{not json")


def test_parse_json_upload_rejects_non_array():
    with pytest.raises(DatasetUploadFormatError, match="JSON array"):
        parse_json_upload(b'{"question":"q"}')


# --- CSV 解析：反向 / 异常 ---


def test_parse_csv_rejects_missing_question_column():
    content = "标准答案,模型答案\nGT,MA\n".encode("utf-8")
    with pytest.raises(DatasetUploadFormatError, match="题目"):
        parse_csv_upload(content)


def test_parse_csv_rejects_header_only():
    content = "题目,模型答案\n".encode("utf-8")
    with pytest.raises(DatasetUploadFormatError, match="无有效数据"):
        parse_csv_upload(content)


def test_parse_csv_rejects_non_utf8():
    with pytest.raises(DatasetUploadFormatError, match="UTF-8"):
        parse_csv_upload(b"\xff\xfe\xfd")


def test_parse_csv_skips_blank_rows():
    content = (
        "题目,模型答案\n"
        "Q1,A1\n"
        "\n"
        "Q2,A2\n"
    ).encode("utf-8")
    items = parse_csv_upload(content)
    assert len(items) == 2


def test_parse_csv_english_headers():
    content = "question,ground_truth,model_answer\nQ,GT,MA\n".encode("utf-8")
    items = parse_csv_upload(content)
    assert items[0]["question"] == "Q"
    assert items[0]["ground_truth"] == "GT"


# --- 校验 ---


def test_validate_json_missing_model_answer():
    # model_answer 现为可选字段（空时由 ground_truth 回退），改为测 question 缺失
    errors = validate_dataset_items_with_format(
        [{"model_answer": "a"}],
        "data.json",
    )
    assert len(errors) == 1
    assert "Item 0" in errors[0]


def test_validate_csv_row_label():
    errors = validate_dataset_items_with_format(
        [{"model_answer": "a"}],
        "data.csv",
    )
    assert len(errors) == 1
    assert "第 2 行" in errors[0]


def test_export_empty_dataset():
    assert export_dataset_json_bytes([]) == b"[]"
    csv_body = export_dataset_csv_bytes([])
    assert "Question" in csv_body.decode("utf-8-sig")


@pytest.fixture
async def export_client():
    engine = create_async_engine("sqlite+aiosqlite:///:memory:")
    async with engine.begin() as conn:
        await conn.run_sync(Base.metadata.create_all)
    factory = async_sessionmaker(engine, class_=AsyncSession, expire_on_commit=False)

    async def _override():
        async with factory() as session:
            yield session

    app.dependency_overrides[get_session] = _override
    async with factory() as session:
        session.add(
            Dataset(
                id="ds_export",
                name="导出测试",
                data=json.dumps(
                    [{"question": "Q", "model_answer": "A", "ground_truth": "G"}],
                    ensure_ascii=False,
                ),
                item_count=1,
            )
        )
        await session.commit()

    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        yield client

    app.dependency_overrides.clear()
    await engine.dispose()


@pytest.mark.asyncio
async def test_export_json_and_csv(export_client):
    j = await export_client.get("/datasets/ds_export/export", params={"format": "json"})
    assert j.status_code == 200
    assert "application/json" in j.headers.get("content-type", "")
    body = json.loads(j.content.decode("utf-8"))
    assert body[0]["question"] == "Q"

    c = await export_client.get("/datasets/ds_export/export", params={"format": "csv"})
    assert c.status_code == 200
    assert c.content.startswith(b"\xef\xbb\xbf")
    assert "Question" in c.content.decode("utf-8-sig")


@pytest.mark.asyncio
async def test_upload_csv_via_api(export_client):
    csv_body = (
        "\ufeff题目,标准答案,模型答案\n"
        "导入题,GT,MA\n"
    ).encode("utf-8")
    res = await export_client.post(
        "/upload",
        files={"file": ("sample.csv", csv_body, "text/csv")},
        data={"name": "csv导入测试"},
    )
    assert res.status_code == 200
    assert res.json()["item_count"] == 1


@pytest.mark.asyncio
async def test_upload_json_via_api(export_client):
    payload = json.dumps([{"question": "JQ", "model_answer": "JA"}]).encode("utf-8")
    res = await export_client.post(
        "/upload",
        files={"file": ("sample.json", payload, "application/json")},
    )
    assert res.status_code == 200
    assert res.json()["item_count"] == 1


@pytest.mark.asyncio
async def test_upload_rejects_unsupported_extension(export_client):
    res = await export_client.post(
        "/upload",
        files={"file": ("bad.txt", b"hello", "text/plain")},
    )
    assert res.status_code == 400


@pytest.mark.asyncio
async def test_upload_rejects_invalid_json(export_client):
    res = await export_client.post(
        "/upload",
        files={"file": ("bad.json", b"not-array", "application/json")},
    )
    assert res.status_code == 400


@pytest.mark.asyncio
async def test_upload_rejects_validation_errors(export_client):
    # question 缺失时应被校验拒绝
    payload = json.dumps([{"model_answer": "only-a"}]).encode("utf-8")
    res = await export_client.post(
        "/upload",
        files={"file": ("bad.json", payload, "application/json")},
    )
    assert res.status_code == 400
    body = res.json()
    assert body["detail"]["error"] == "Validation failed"


@pytest.mark.asyncio
async def test_download_csv_template(export_client):
    res = await export_client.get("/datasets/template.csv")
    assert res.status_code == 200
    assert res.content.startswith(b"\xef\xbb\xbf")
    assert "示例题目" in res.content.decode("utf-8-sig")


@pytest.mark.asyncio
async def test_export_not_found(export_client):
    res = await export_client.get("/datasets/ds_missing/export")
    assert res.status_code == 404


@pytest.mark.asyncio
async def test_export_invalid_format_returns_422(export_client):
    res = await export_client.get(
        "/datasets/ds_export/export",
        params={"format": "xml"},
    )
    assert res.status_code == 422


@pytest.mark.asyncio
async def test_export_reimport_roundtrip(export_client):
    exported = await export_client.get(
        "/datasets/ds_export/export",
        params={"format": "json"},
    )
    assert exported.status_code == 200
    res = await export_client.post(
        "/upload",
        files={
            "file": (
                "roundtrip.json",
                exported.content,
                "application/json",
            )
        },
        data={"name": "roundtrip"},
    )
    assert res.status_code == 200
    assert res.json()["item_count"] == 1
    new_id = res.json()["dataset_id"]
    assert new_id != "ds_export"


# ---------------------------------------------------------------------------
# 场景子步骤导出（Step_* 列还原 / round-trip）
# ---------------------------------------------------------------------------


def _make_scenario_items() -> list:
    """构造带 steps 的 items（模拟 load_steps_for_dataset_items 挂载后的结构）。"""
    return [
        {
            "id": "ditem_1",
            "question": "示例题目：卖点分析",
            "prompt": "你是数据质量评测专家",
            "ground_truth": "GT 主",
            "model_answer": "MA 主",
            "judging_mode": "Ground_Truth评测",
            "steps": [
                {
                    "name": "步骤1：客户声音样本",
                    "weight": 0.3,
                    "enabled": True,
                    "gt_expected": {
                        "keywords": ["样本量", "数据来源"],
                        "text": "样本量,数据来源",
                        "model_answer": "MA 步骤1",
                    },
                    "match_rules": {"eval_config": {"mode": "keyword"}},
                },
                {
                    "name": "步骤2：竞品定位",
                    "weight": 0.7,
                    "enabled": True,
                    "gt_expected": {"keywords": [], "model_answer": "MA 步骤2"},
                    "match_rules": {
                        "eval_config": {
                            "mode": "llm_aggregate",
                            "aggregate_prompt": "## 角色\n你是评测专家",
                        }
                    },
                },
            ],
        }
    ]


def test_step_to_export_fields_keyword_step():
    step = {
        "name": "步骤1",
        "weight": 0.3,
        "gt_expected": {"keywords": ["样本量", "数据来源"], "text": "样本量,数据来源"},
        "match_rules": {"eval_config": {"mode": "keyword"}},
    }
    fields = _step_to_export_fields(step)
    assert fields["Step_Question"] == "步骤1"
    assert fields["Step_Judging_Mode"] == "Ground_Truth评测"
    assert fields["Step_Prompt"] == ""  # keyword 模式无 prompt
    assert fields["Step_Ground_Truth"] == "样本量,数据来源"
    assert fields["Step_Mode_Answer"] == ""  # 无 model_answer
    assert fields["Step_Weight"] == "0.3"


def test_step_to_export_fields_llm_aggregate_step_preserves_prompt():
    step = {
        "name": "步骤2",
        "weight": 0.7,
        "gt_expected": {"keywords": [], "model_answer": "MA2"},
        "match_rules": {
            "eval_config": {
                "mode": "llm_aggregate",
                "aggregate_prompt": "你是评测专家",
            }
        },
    }
    fields = _step_to_export_fields(step)
    assert fields["Step_Judging_Mode"] == "基于Prompt评测"
    assert fields["Step_Prompt"] == "你是评测专家"  # 逆映射 aggregate_prompt
    assert fields["Step_Mode_Answer"] == "MA2"  # 逆映射 gt_expected.model_answer
    assert fields["Step_Ground_Truth"] == ""  # 无 keywords/text


def test_step_to_export_fields_keywords_join_when_no_text():
    step = {
        "name": "步骤3",
        "weight": 0.5,
        "gt_expected": {"keywords": ["A", "B", "C"]},
        "match_rules": {"eval_config": {"mode": "keyword"}},
    }
    fields = _step_to_export_fields(step)
    assert fields["Step_Ground_Truth"] == "A,B,C"


def test_export_csv_with_steps_has_step_columns_and_multi_row():
    import csv as _csv

    items = _make_scenario_items()
    body = export_dataset_csv_bytes(items)
    text = body.decode("utf-8-sig")
    rows = list(_csv.DictReader(io.StringIO(text)))
    # 表头含 Step_* 列
    assert "Step_Question" in rows[0]
    assert "Step_Mode_Answer" in rows[0]
    # 1 个 item × 2 步 → 2 数据行
    assert len(rows) == 2
    # 首行含 Question，次行 Question 留空
    assert rows[0]["Question"] == "示例题目：卖点分析"
    assert rows[1]["Question"] == ""
    # 次行仍含 Step_Question
    assert rows[1]["Step_Question"] == "步骤2：竞品定位"


def test_export_csv_without_steps_has_no_step_columns():
    items = [{"id": "d1", "question": "Q", "model_answer": "A"}]
    text = export_dataset_csv_bytes(items).decode("utf-8-sig")
    headers = text.splitlines()[0].split(",")
    assert "Step_Question" not in headers
    assert "Question" in headers


def test_scenario_csv_export_roundtrips_through_parse_scenario_csv():
    items = _make_scenario_items()
    body = export_dataset_csv_bytes(items)
    parsed = parse_scenario_csv(body)
    # 1 题 2 步
    assert len(parsed["items"]) == 1
    assert parsed["steps_count"] == 2
    item = parsed["items"][0]
    assert item["question"] == "示例题目：卖点分析"
    assert item["model_answer"] == "MA 主"
    steps = parsed["scenario_steps"]
    # 步骤逆映射还原
    s1, s2 = steps
    assert s1["step_question"] == "步骤1：客户声音样本"
    assert s1["step_ground_truth"] == "样本量,数据来源"
    assert s1["step_model_answer"] == "MA 步骤1"
    assert s1["step_judging_mode"] == "Ground_Truth评测"
    assert s2["step_question"] == "步骤2：竞品定位"
    assert s2["step_judging_mode"] == "基于Prompt评测"
    assert s2["step_prompt"] == "## 角色\n你是评测专家"  # aggregate_prompt 还原
    assert s2["step_model_answer"] == "MA 步骤2"


def test_export_xlsx_with_steps_has_step_columns():
    items = _make_scenario_items()
    body = export_dataset_xlsx_bytes(items)
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(body), read_only=True, data_only=True)
    ws = wb.active
    header_row = [str(c) for c in next(ws.iter_rows(values_only=True))]
    wb.close()
    assert "Step_Question" in header_row
    assert "Step_Mode_Answer" in header_row


def test_export_json_preserves_steps_field():
    items = _make_scenario_items()
    body = export_dataset_json_bytes(items)
    data = json.loads(body)
    assert data and "steps" in data[0]
    assert data[0]["steps"][0]["name"] == "步骤1：客户声音样本"


@pytest.fixture
async def scenario_export_client():
    """创建带子步骤模板的评测集（含 disabled 步骤 + llm_aggregate 步骤）。"""
    from services.dataset_item_scenario_service import upsert_item_template

    engine = create_async_engine("sqlite+aiosqlite:///:memory:")
    async with engine.begin() as conn:
        await conn.run_sync(Base.metadata.create_all)
    factory = async_sessionmaker(engine, class_=AsyncSession, expire_on_commit=False)

    async def _override():
        async with factory() as session:
            yield session

    app.dependency_overrides[get_session] = _override
    dataset_item_id = "ditem_1"
    async with factory() as session:
        session.add(
            Dataset(
                id="ds_sc_export",
                name="场景导出测试",
                data=json.dumps(
                    [
                        {
                            "id": dataset_item_id,
                            "question": "场景题",
                            "ground_truth": "GT",
                            "model_answer": "MA",
                        }
                    ],
                    ensure_ascii=False,
                ),
                item_count=1,
            )
        )
        await session.flush()
        await upsert_item_template(
            session,
            dataset_id="ds_sc_export",
            dataset_item_id=dataset_item_id,
            name="Q1 场景题",
            steps=[
                {
                    "step_key": "S1",
                    "name": "关键词步",
                    "step_order": 1,
                    "weight": 0.4,
                    "enabled": True,
                    "gt_expected": {
                        "keywords": ["样本量", "数据来源"],
                        "text": "样本量,数据来源",
                    },
                    "match_rules": {"eval_config": {"mode": "keyword"}},
                },
                {
                    "step_key": "S2",
                    "name": "提示词步",
                    "step_order": 2,
                    "weight": 0.6,
                    "enabled": True,
                    "gt_expected": {"keywords": [], "model_answer": "MA2"},
                    "match_rules": {
                        "eval_config": {
                            "mode": "llm_aggregate",
                            "aggregate_prompt": "你是评测专家",
                        }
                    },
                },
                {
                    "step_key": "S3",
                    "name": "禁用步",
                    "step_order": 3,
                    "weight": 0.0,
                    "enabled": False,
                    "gt_expected": {"keywords": ["禁用"]},
                    "match_rules": {"eval_config": {"mode": "keyword"}},
                },
            ],
        )
        await session.commit()

    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        yield client

    app.dependency_overrides.clear()
    await engine.dispose()


@pytest.mark.asyncio
async def test_export_route_csv_includes_enabled_steps_only(scenario_export_client):
    import csv as _csv

    res = await scenario_export_client.get(
        "/datasets/ds_sc_export/export", params={"format": "csv"}
    )
    assert res.status_code == 200
    text = res.content.decode("utf-8-sig")
    rows = list(_csv.DictReader(io.StringIO(text)))
    assert "Step_Question" in rows[0]
    # 仅 2 个 enabled 步骤 → 2 数据行（disabled 步骤不导出）
    assert len(rows) == 2
    names = [r["Step_Question"] for r in rows]
    assert "禁用步" not in names
    assert "关键词步" in names
    assert "提示词步" in names
    # llm_aggregate 的 aggregate_prompt 还原为 Step_Prompt
    prompt_row = next(r for r in rows if r["Step_Question"] == "提示词步")
    assert prompt_row["Step_Prompt"] == "你是评测专家"


@pytest.mark.asyncio
async def test_export_route_json_includes_steps(scenario_export_client):
    res = await scenario_export_client.get(
        "/datasets/ds_sc_export/export", params={"format": "json"}
    )
    assert res.status_code == 200
    data = json.loads(res.content.decode("utf-8"))
    assert "steps" in data[0]
    # 仅 enabled 步骤
    assert len(data[0]["steps"]) == 2
    names = [s["name"] for s in data[0]["steps"]]
    assert "禁用步" not in names


@pytest.mark.asyncio
async def test_export_route_csv_without_steps_unchanged(export_client):
    """无子步骤的评测集导出 CSV → 不含 Step_* 列（回归）。"""
    res = await export_client.get(
        "/datasets/ds_export/export", params={"format": "csv"}
    )
    assert res.status_code == 200
    headers = res.content.decode("utf-8-sig").splitlines()[0].split(",")
    assert "Step_Question" not in headers
    assert "Question" in headers


@pytest.fixture
async def multi_item_scenario_client():
    """2 个 item，每个都有模板；item2 含一个禁用步骤。"""
    from services.dataset_item_scenario_service import upsert_item_template

    engine = create_async_engine("sqlite+aiosqlite:///:memory:")
    async with engine.begin() as conn:
        await conn.run_sync(Base.metadata.create_all)
    factory = async_sessionmaker(engine, class_=AsyncSession, expire_on_commit=False)

    async def _override():
        async with factory() as session:
            yield session

    app.dependency_overrides[get_session] = _override
    async with factory() as session:
        session.add(
            Dataset(
                id="ds_multi",
                name="多题场景",
                data=json.dumps(
                    [
                        {"id": "ditem_1", "question": "题一", "model_answer": "A1"},
                        {"id": "ditem_2", "question": "题二", "model_answer": "A2"},
                    ],
                    ensure_ascii=False,
                ),
                item_count=2,
            )
        )
        await session.flush()
        # item1: 2 个全启用步骤
        await upsert_item_template(
            session,
            dataset_id="ds_multi",
            dataset_item_id="ditem_1",
            name="Q1",
            steps=[
                {
                    "step_key": "S1",
                    "name": "题一步骤1",
                    "step_order": 1,
                    "weight": 0.5,
                    "enabled": True,
                    "gt_expected": {"keywords": ["k1"]},
                    "match_rules": {"eval_config": {"mode": "keyword"}},
                },
                {
                    "step_key": "S2",
                    "name": "题一步骤2",
                    "step_order": 2,
                    "weight": 0.5,
                    "enabled": True,
                    "gt_expected": {"keywords": ["k2"]},
                    "match_rules": {"eval_config": {"mode": "keyword"}},
                },
            ],
        )
        # item2: 3 个步骤，其中 S3 禁用
        await upsert_item_template(
            session,
            dataset_id="ds_multi",
            dataset_item_id="ditem_2",
            name="Q2",
            steps=[
                {
                    "step_key": "S1",
                    "name": "题二步骤1",
                    "step_order": 1,
                    "weight": 0.4,
                    "enabled": True,
                    "gt_expected": {"keywords": ["k3"]},
                    "match_rules": {"eval_config": {"mode": "keyword"}},
                },
                {
                    "step_key": "S2",
                    "name": "题二步骤2",
                    "step_order": 2,
                    "weight": 0.4,
                    "enabled": True,
                    "gt_expected": {"keywords": ["k4"]},
                    "match_rules": {"eval_config": {"mode": "keyword"}},
                },
                {
                    "step_key": "S3",
                    "name": "题二禁用步",
                    "step_order": 3,
                    "weight": 0.2,
                    "enabled": False,
                    "gt_expected": {"keywords": ["k5"]},
                    "match_rules": {"eval_config": {"mode": "keyword"}},
                },
            ],
        )
        await session.commit()

    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        yield client

    app.dependency_overrides.clear()
    await engine.dispose()


@pytest.mark.asyncio
async def test_multi_item_csv_exports_all_items_steps(multi_item_scenario_client):
    """确认多 item 场景下每个 item 的 enabled 子步骤都导出。"""
    import csv as _csv

    res = await multi_item_scenario_client.get(
        "/datasets/ds_multi/export", params={"format": "csv"}
    )
    assert res.status_code == 200
    rows = list(_csv.DictReader(io.StringIO(res.content.decode("utf-8-sig"))))
    # item1: 2 步 + item2: 2 步（禁用步不导出）= 4 行
    assert len(rows) == 4
    questions = [r["Question"] for r in rows]
    # 第 1 行是题一，第 3 行是题二
    assert questions[0] == "题一"
    assert questions[2] == "题二"
    # 题一的两个步骤都在
    step_qs = [r["Step_Question"] for r in rows]
    assert "题一步骤1" in step_qs
    assert "题一步骤2" in step_qs
    assert "题二步骤1" in step_qs
    assert "题二步骤2" in step_qs
    # 禁用步不在
    assert "题二禁用步" not in step_qs


@pytest.mark.asyncio
async def test_multi_item_json_exports_all_items_steps(multi_item_scenario_client):
    res = await multi_item_scenario_client.get(
        "/datasets/ds_multi/export", params={"format": "json"}
    )
    assert res.status_code == 200
    data = json.loads(res.content.decode("utf-8"))
    assert len(data) == 2
    # 两个 item 都有 steps
    assert len(data[0]["steps"]) == 2
    assert len(data[1]["steps"]) == 2  # 禁用步不导出 → 2 个 enabled
    assert "题二禁用步" not in [s["name"] for s in data[1]["steps"]]


@pytest.mark.asyncio
async def test_scenario_import_then_export_roundtrip_via_api(export_client):
    """端到端：通过 API 导入 2 题场景 CSV → 导出 CSV → 确认所有题的步骤都还原。"""
    import csv as _csv

    csv_body = (
        "\ufeffQuestion,Prompt,Ground_Truth,Mode_Answer,Judging_Mode,"
        "Step_Question,Step_Prompt,Step_Ground_Truth,Step_Mode_Answer,"
        "Step_Weight,Step_Judging_Mode\n"
        "题一,P1,GT1,MA1,Ground_Truth评测,题一步1,,kw1,,0.5,Ground_Truth评测\n"
        ",,,,,题一步2,,kw2,,0.5,Ground_Truth评测\n"
        "题二,P2,GT2,MA2,Ground_Truth评测,题二步1,,kw3,,0.5,Ground_Truth评测\n"
        ",,,,,题二步2,,kw4,,0.5,Ground_Truth评测\n"
    ).encode("utf-8")
    res = await export_client.post(
        "/datasets/import/with-scenario",
        files={"file": ("scenario.csv", csv_body, "text/csv")},
    )
    assert res.status_code == 200
    dataset_id = res.json()["dataset_id"]

    # 导出 CSV
    res = await export_client.get(
        f"/datasets/{dataset_id}/export", params={"format": "csv"}
    )
    assert res.status_code == 200
    rows = list(_csv.DictReader(io.StringIO(res.content.decode("utf-8-sig"))))
    # 2 题 × 2 步 = 4 行
    assert len(rows) == 4
    step_qs = [r["Step_Question"] for r in rows]
    assert "题一步1" in step_qs
    assert "题一步2" in step_qs
    assert "题二步1" in step_qs
    assert "题二步2" in step_qs
    # 两题的 Question 都出现在首行
    questions = [r["Question"] for r in rows if r["Question"]]
    assert "题一" in questions
    assert "题二" in questions


# ---------------------------------------------------------------------------
# 导入重名 → 自动追加 _2/_3 编号
# ---------------------------------------------------------------------------


@pytest.mark.asyncio
async def test_upload_duplicate_name_appends_suffix(export_client):
    """同名评测集二次导入时，自动加 _2 后缀，不报错也不覆盖。"""
    payload = json.dumps([{"question": "Q", "model_answer": "A"}]).encode("utf-8")
    first = await export_client.post(
        "/upload",
        files={"file": ("dup.json", payload, "application/json")},
        data={"name": "重名测试集"},
    )
    assert first.status_code == 200
    assert first.json()["item_count"] == 1

    # 第二次同名导入 → 应自动改名为「重名测试集_2」
    second = await export_client.post(
        "/upload",
        files={"file": ("dup.json", payload, "application/json")},
        data={"name": "重名测试集"},
    )
    assert second.status_code == 200
    assert second.json()["item_count"] == 1

    # 第三次 → 「重名测试集_3」
    third = await export_client.post(
        "/upload",
        files={"file": ("dup.json", payload, "application/json")},
        data={"name": "重名测试集"},
    )
    assert third.status_code == 200

    # 拉取评测集列表，确认三套都在且名称分别为原名/_2/_3
    lst = await export_client.get("/datasets")
    assert lst.status_code == 200
    names = [d["name"] for d in lst.json()["datasets"]]
    assert "重名测试集" in names
    assert "重名测试集_2" in names
    assert "重名测试集_3" in names


@pytest.mark.asyncio
async def test_scenario_import_duplicate_name_appends_suffix(export_client):
    """场景导入同名时也应自动编号。"""
    csv_body = (
        "\ufeffQuestion,Prompt,Ground_Truth,Mode_Answer,Judging_Mode,"
        "Step_Question,Step_Prompt,Step_Ground_Truth,Step_Mode_Answer,"
        "Step_Weight,Step_Judging_Mode\n"
        "题一,P1,GT1,MA1,Ground_Truth评测,步1,,kw1,,1,Ground_Truth评测\n"
    ).encode("utf-8")
    first = await export_client.post(
        "/datasets/import/with-scenario",
        files={"file": ("sc.csv", csv_body, "text/csv")},
        data={"name": "场景重名"},
    )
    assert first.status_code == 200

    second = await export_client.post(
        "/datasets/import/with-scenario",
        files={"file": ("sc.csv", csv_body, "text/csv")},
        data={"name": "场景重名"},
    )
    assert second.status_code == 200

    lst = await export_client.get("/datasets")
    names = [d["name"] for d in lst.json()["datasets"]]
    assert "场景重名" in names
    assert "场景重名_2" in names
